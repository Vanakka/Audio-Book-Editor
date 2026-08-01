"""Parse a Libation Library Export XLSX file for audiobook metadata."""

import logging
from pathlib import Path

from core.models import MetadataResult

logger = logging.getLogger(__name__)


def parse_libation_export(xlsx_path: str) -> dict[str, MetadataResult]:
    """Parse a Libation XLSX export. Returns dict keyed by ASIN."""
    try:
        from openpyxl import load_workbook
    except ImportError:
        logger.error("openpyxl not installed, cannot parse Libation export")
        return {}

    path = Path(xlsx_path)
    if not path.exists():
        return {}

    results = {}
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active

        # Find header row and column indices
        headers = {}
        for row in ws.iter_rows(min_row=1, max_row=1, values_only=False):
            for cell in row:
                if cell.value:
                    headers[str(cell.value).strip()] = cell.column - 1  # 0-indexed

        if not headers:
            wb.close()
            return {}

        # Map expected column names
        col_map = {
            "asin": _find_col(headers, "Audible Product Id", "AudibleProductId", "Asin", "ASIN"),
            "title": _find_col(headers, "Title"),
            "subtitle": _find_col(headers, "Subtitle"),
            "authors": _find_col(headers, "Authors", "Author"),
            "narrators": _find_col(headers, "Narrators", "Narrator"),
            "description": _find_col(headers, "Description"),
            "publisher": _find_col(headers, "Publisher"),
            "series_names": _find_col(headers, "Series Names", "SeriesNames", "Series"),
            "series_order": _find_col(headers, "Series Order", "SeriesOrder"),
            "cover_id_large": _find_col(headers, "Cover Id Large", "CoverIdLarge"),
            "cover_id": _find_col(headers, "Cover Id", "CoverId"),
            "date_published": _find_col(headers, "Date Published", "DatePublished"),
            "categories": _find_col(headers, "Categories", "Category"),
        }

        asin_col = col_map.get("asin")
        if asin_col is None:
            wb.close()
            return {}

        for row in ws.iter_rows(min_row=2, values_only=True):
            row = list(row)
            asin = _get_cell(row, asin_col)
            if not asin:
                continue

            # Parse series info
            series_name = ""
            series_number = ""
            series_names_val = _get_cell(row, col_map.get("series_names"))
            series_order_val = _get_cell(row, col_map.get("series_order"))

            if series_names_val:
                series_name = series_names_val.split(",")[0].strip()
            if series_order_val:
                # Format can be "5 : Series Name" or just "5"
                parts = str(series_order_val).split(":")
                series_number = parts[0].strip()

            # Build cover URL from Amazon image ID
            cover_url = ""
            cover_id = _get_cell(row, col_map.get("cover_id_large")) or _get_cell(row, col_map.get("cover_id"))
            if cover_id:
                cover_url = f"https://m.media-amazon.com/images/I/{cover_id}._SL2400_.jpg"

            # Parse year from date
            year = ""
            date_val = _get_cell(row, col_map.get("date_published"))
            if date_val:
                year = str(date_val)[:4]

            # Clean description (may contain HTML)
            description = _get_cell(row, col_map.get("description"))
            if description:
                description = _strip_html(description)

            result = MetadataResult(
                source="libation",
                title=_get_cell(row, col_map.get("title")),
                subtitle=_get_cell(row, col_map.get("subtitle")),
                author=_get_cell(row, col_map.get("authors")),
                narrator=_get_cell(row, col_map.get("narrators")),
                series=series_name,
                series_number=series_number,
                description=description,
                publisher=_get_cell(row, col_map.get("publisher")),
                year=year,
                genre=_get_cell(row, col_map.get("categories")),
                cover_url=cover_url,
            )
            results[asin] = result

        wb.close()
    except Exception as e:
        logger.error("Error parsing Libation export: %s", e)

    return results


def _find_col(headers: dict, *names: str) -> int | None:
    for name in names:
        if name in headers:
            return headers[name]
        # Case-insensitive fallback
        for h, idx in headers.items():
            if h.lower() == name.lower():
                return idx
    return None


def _get_cell(row: list, col: int | None) -> str:
    if col is None or col >= len(row):
        return ""
    val = row[col]
    if val is None:
        return ""
    return str(val).strip()


def _strip_html(text: str) -> str:
    """Remove basic HTML tags from text."""
    import re
    text = re.sub(r'<br\s*/?>', '\n', text, flags=re.IGNORECASE)
    text = re.sub(r'<p\s*/?>', '\n', text, flags=re.IGNORECASE)
    text = re.sub(r'</p>', '\n', text, flags=re.IGNORECASE)
    text = re.sub(r'<[^>]+>', '', text)
    text = re.sub(r'\n{3,}', '\n\n', text)
    return text.strip()
